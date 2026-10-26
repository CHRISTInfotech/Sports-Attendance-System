from io import BytesIO
from datetime import datetime
import os
import tempfile

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.test import override_settings
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from unittest.mock import patch
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .models import AttendanceRecord, EmailNotificationJob, Meeting, Membership, School, Session, Sport, Team, UserProfile, Venue
from .views import (
    _management_bulk_workbook,
    _simple_status_bulk_workbook,
    build_sessions_bulk_upload_workbook,
    import_schools_from_file,
    import_sessions_from_file,
    import_sports_from_file,
    import_students_from_file,
    import_teams_from_file,
    import_trainers_from_file,
    import_venues_from_file,
    read_bulk_upload_rows,
)
from .email_notifications import queue_notification_emails, send_notification_emails


User = get_user_model()


@override_settings(
    EMAIL_NOTIFICATIONS_ENABLED=True,
    EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
    DEFAULT_FROM_EMAIL="sports.lavasa@christuniversity.in",
)
class EmailNotificationTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="mail-recipient@christuniversity.in",
            email="mail-recipient@christuniversity.in",
            first_name="Mail",
            last_name="Recipient",
        )

    def test_sends_one_html_email_for_duplicate_recipient(self):
        send_notification_emails(
            [self.user, self.user],
            "meeting_scheduled",
            {"title": "New meeting scheduled", "message": "A meeting was scheduled.", "details": [("Venue", "Seminar Hall")]},
        )

        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, [self.user.email])
        self.assertIn("Meeting Scheduled", mail.outbox[0].subject)
        self.assertTrue(mail.outbox[0].alternatives)
        html = mail.outbox[0].alternatives[0].content
        self.assertIn("Dear Mail Recipient", html)
        self.assertIn("Seminar Hall", html)
        self.assertNotIn("{{", html)
        self.assertNotIn("}}", html)

    @patch("core.email_notifications.get_connection")
    def test_queue_does_not_open_smtp_connection(self, mocked_get_connection):
        queue_notification_emails(
            [self.user, self.user],
            "practice_scheduled",
            {"title": "Practice", "message": "Practice scheduled."},
        )

        mocked_get_connection.assert_not_called()
        job = EmailNotificationJob.objects.get()
        self.assertEqual(job.recipient_user_ids, [self.user.pk])

    def test_worker_delivers_queued_email(self):
        job = queue_notification_emails(
            [self.user],
            "practice_scheduled",
            {"title": "Practice", "message": "Practice scheduled."},
        )

        call_command("run_email_worker", once=True, verbosity=0)

        job.refresh_from_db()
        self.assertEqual(job.status, EmailNotificationJob.Status.SENT)
        self.assertEqual(len(mail.outbox), 1)

    @patch("core.email_notifications.EmailMultiAlternatives.send", side_effect=RuntimeError("SMTP unavailable"))
    def test_delivery_failure_does_not_escape(self, mocked_send):
        send_notification_emails(
            [self.user],
            "feedback_received",
            {"title": "New feedback", "message": "Feedback received."},
        )
        mocked_send.assert_called_once()

    @patch("core.email_notifications.get_connection")
    def test_batch_authentication_failure_is_attempted_only_once(self, mocked_get_connection):
        second_user = User.objects.create_user(
            username="second-recipient@christuniversity.in",
            email="second-recipient@christuniversity.in",
        )
        connection = mocked_get_connection.return_value
        connection.open.side_effect = RuntimeError("Authentication rejected")

        send_notification_emails(
            [self.user, second_user],
            "practice_scheduled",
            {"title": "Practice", "message": "Practice scheduled."},
        )

        connection.open.assert_called_once_with()


class TrainerStudentPermissionTests(TestCase):
    def setUp(self):
        self.trainer = self.make_user("trainer@christuniversity.in", UserProfile.Role.TRAINER)
        self.other_trainer = self.make_user("other-trainer@christuniversity.in", UserProfile.Role.TRAINER)
        self.assigned_student = self.make_user("assigned@christuniversity.in", UserProfile.Role.MEMBER)
        self.other_student = self.make_user("other@christuniversity.in", UserProfile.Role.MEMBER)
        self.sport = Sport.objects.create(name="Cricket")
        self.other_sport = Sport.objects.create(name="Football")
        self.assigned_team = Team.objects.create(
            sport=self.sport,
            name="University Team",
            team_type=Team.TeamType.UNIVERSITY,
            gender=Team.TeamGender.MALE,
            coordinator=self.trainer,
        )
        self.other_team = Team.objects.create(
            sport=self.other_sport,
            name="Reserve Team",
            team_type=Team.TeamType.RESERVE,
            gender=Team.TeamGender.FEMALE,
            coordinator=self.other_trainer,
        )
        self.venue = Venue.objects.create(name="Ground A", location="Sports block")
        self.assigned_membership = Membership.objects.create(user=self.assigned_student, team=self.assigned_team)
        self.other_membership = Membership.objects.create(user=self.other_student, team=self.other_team)

    def make_user(self, email, role):
        user = User.objects.create_user(username=email, email=email, password="pass")
        profile, _ = UserProfile.objects.get_or_create(user=user)
        profile.role = role
        profile.save(update_fields=["role", "updated_at"])
        return user

    def test_trainer_sees_only_students_from_assigned_teams(self):
        self.client.force_login(self.trainer)

        response = self.client.get(reverse("members"))

        self.assertContains(response, self.assigned_student.email)
        self.assertNotContains(response, self.other_student.email)

    def test_trainer_can_add_student_only_to_assigned_team(self):
        self.client.force_login(self.trainer)

        response = self.client.post(reverse("members"), {
            "action": "create",
            "student_name": "New Player",
            "student_email": "new.player@christuniversity.in",
            "gender": "Female",
            "sport": self.assigned_team.sport_id,
            "team_type": self.assigned_team.team_type,
            "team": self.assigned_team.pk,
            "is_active": "on",
        })

        self.assertRedirects(response, reverse("members"))
        self.assertTrue(Membership.objects.filter(user__email="new.player@christuniversity.in", team=self.assigned_team).exists())
        self.assertEqual(User.objects.get(email="new.player@christuniversity.in").profile.gender, "Female")

    def test_trainer_cannot_add_student_to_other_team(self):
        self.client.force_login(self.trainer)

        response = self.client.post(reverse("members"), {
            "action": "create",
            "student_name": "Blocked Player",
            "student_email": "blocked.player@christuniversity.in",
            "team": self.other_team.pk,
            "is_active": "on",
        })

        self.assertRedirects(response, reverse("members"))
        self.assertFalse(Membership.objects.filter(user__email="blocked.player@christuniversity.in", team=self.other_team).exists())

    def test_trainer_can_deactivate_and_delete_assigned_membership(self):
        self.client.force_login(self.trainer)

        self.client.post(reverse("members"), {
            "action": "deactivate",
            "membership_id": self.assigned_membership.pk,
        })
        self.assigned_membership.refresh_from_db()
        self.assertFalse(self.assigned_membership.is_active)

        self.client.post(reverse("members"), {
            "action": "delete",
            "membership_id": self.assigned_membership.pk,
        })
        self.assertFalse(Membership.objects.filter(pk=self.assigned_membership.pk).exists())

    def test_trainer_cannot_modify_other_team_membership(self):
        self.client.force_login(self.trainer)

        response = self.client.post(reverse("members"), {
            "action": "deactivate",
            "membership_id": self.other_membership.pk,
        })

        self.assertRedirects(response, reverse("members"))
        self.other_membership.refresh_from_db()
        self.assertTrue(self.other_membership.is_active)

    def test_trainer_student_detail_is_restricted_to_assigned_team_records(self):
        session = Session.objects.create(
            team=self.assigned_team,
            title="Practice session",
            start_at=timezone.now(),
            end_at=timezone.now(),
            venue="Ground A",
            scheduled_by=self.trainer,
        )
        other_session = Session.objects.create(
            team=self.other_team,
            title="Other practice",
            start_at=timezone.now(),
            end_at=timezone.now(),
            venue="Ground B",
            scheduled_by=self.other_trainer,
        )
        AttendanceRecord.objects.create(session=session, member=self.assigned_student, status=AttendanceRecord.Status.PRESENT, marked_by=self.trainer)
        AttendanceRecord.objects.create(session=other_session, member=self.assigned_student, status=AttendanceRecord.Status.ABSENT, marked_by=self.other_trainer)
        Membership.objects.create(user=self.assigned_student, team=self.other_team)
        self.client.force_login(self.trainer)

        response = self.client.get(reverse("member_detail", args=[self.assigned_student.pk]))

        self.assertContains(response, "Practice session")
        self.assertContains(response, "100.0%")
        self.assertNotContains(response, "Other practice")

    def test_trainer_can_submit_attendance_for_assigned_team_session(self):
        session = Session.objects.create(
            team=self.assigned_team,
            title="Attendance Practice",
            start_at=timezone.now(),
            end_at=timezone.now(),
            venue="Ground A",
            scheduled_by=self.trainer,
        )
        self.client.force_login(self.trainer)

        response = self.client.post(reverse("take_attendance", args=[session.pk]), {
            f"status_{self.assigned_student.pk}": AttendanceRecord.Status.PRESENT,
            f"remarks_{self.assigned_student.pk}": "On time",
        })

        self.assertRedirects(response, reverse("attendance_detail", args=[session.pk]))
        session.refresh_from_db()
        self.assertTrue(session.attendance_submitted)
        self.assertEqual(session.submitted_by, self.trainer)
        record = AttendanceRecord.objects.get(session=session, member=self.assigned_student)
        self.assertEqual(record.status, AttendanceRecord.Status.PRESENT)
        self.assertEqual(record.remarks, "On time")
        self.assertEqual(record.session.team, self.assigned_team)
        self.assertEqual(record.session.team.sport, self.sport)

    def make_session(self, title, team, start_date, end_date=None, status=Session.Status.SCHEDULED):
        end_date = end_date or start_date
        return Session.objects.create(
            team=team,
            title=title,
            start_at=timezone.make_aware(datetime.combine(start_date, datetime.min.time().replace(hour=6))),
            end_at=timezone.make_aware(datetime.combine(end_date, datetime.min.time().replace(hour=8))),
            venue=self.venue.name,
            scheduled_by=team.coordinator,
            status=status,
        )

    def test_practice_session_filters_work_independently_and_together(self):
        from datetime import date

        early = self.make_session("Early Cricket", self.assigned_team, date(2026, 1, 5), date(2026, 1, 6))
        late = self.make_session("Late Cricket", self.assigned_team, date(2026, 1, 10), date(2026, 1, 12))
        football = self.make_session("Football Practice", self.other_team, date(2026, 1, 10), date(2026, 1, 10))
        admin = self.make_user("admin@christuniversity.in", UserProfile.Role.SUPER_ADMIN)
        self.client.force_login(admin)

        response = self.client.get(reverse("sessions"), {"sport": self.sport.pk})
        self.assertEqual({item.pk for item in response.context["sessions"]}, {early.pk, late.pk})

        response = self.client.get(reverse("sessions"), {"team": self.other_team.pk})
        self.assertEqual([item.pk for item in response.context["sessions"]], [football.pk])

        response = self.client.get(reverse("sessions"), {"start_date": "2026-01-10"})
        self.assertEqual({item.pk for item in response.context["sessions"]}, {late.pk, football.pk})

        response = self.client.get(reverse("sessions"), {"end_date": "2026-01-06"})
        self.assertEqual([item.pk for item in response.context["sessions"]], [early.pk])

        response = self.client.get(reverse("sessions"), {
            "sport": self.sport.pk,
            "team": self.assigned_team.pk,
            "start_date": "2026-01-05",
            "end_date": "2026-01-06",
        })
        self.assertEqual([item.pk for item in response.context["sessions"]], [early.pk])
        self.assertEqual(response.context["form"]["sport"].value(), str(self.sport.pk))
        self.assertEqual(response.context["form"]["team"].value(), str(self.assigned_team.pk))
        self.assertEqual(response.context["form"]["start_date"].value(), "2026-01-05")
        self.assertEqual(response.context["form"]["end_date"].value(), "2026-01-06")

    def test_cancelled_session_can_be_moved_to_deleted_sessions_and_viewed(self):
        from datetime import date

        session = self.make_session("Cancelled Practice", self.assigned_team, date(2026, 2, 1), status=Session.Status.CANCELLED)
        self.client.force_login(self.trainer)

        response = self.client.post(reverse("sessions"), {"action": "archive", "session_id": session.pk})

        self.assertRedirects(response, reverse("sessions"))
        self.assertFalse(Session.objects.filter(pk=session.pk).exists())
        deleted = Session.all_objects.get(pk=session.pk)
        self.assertTrue(deleted.is_deleted)
        self.assertEqual(deleted.deleted_by, self.trainer)
        self.assertIsNotNone(deleted.deleted_at)

        response = self.client.get(reverse("deleted_sessions"))
        self.assertContains(response, "Cancelled Practice")
        self.assertContains(response, "Deleted Session Details")
        self.assertContains(response, self.assigned_team.name)

    def test_non_cancelled_session_cannot_be_moved_to_deleted_sessions(self):
        from datetime import date

        session = self.make_session("Active Practice", self.assigned_team, date(2026, 2, 2))
        self.client.force_login(self.trainer)

        self.client.post(reverse("sessions"), {"action": "archive", "session_id": session.pk})

        self.assertTrue(Session.objects.filter(pk=session.pk).exists())
        self.assertFalse(Session.all_objects.get(pk=session.pk).is_deleted)

    def test_completed_session_cannot_be_cancelled_and_has_no_cancel_action(self):
        from datetime import date

        session = self.make_session("Completed Practice", self.assigned_team, date(2026, 2, 3))
        session.attendance_submitted = True
        session.save(update_fields=["attendance_submitted"])
        self.client.force_login(self.trainer)

        response = self.client.get(reverse("sessions"))
        self.assertNotContains(response, "Cancel Session")

        self.client.post(reverse("sessions"), {"action": "cancel", "session_id": session.pk})
        session.refresh_from_db()
        self.assertEqual(session.status, Session.Status.SCHEDULED)

    def test_cancelled_meeting_can_be_moved_to_deleted_meetings_and_viewed(self):
        from datetime import date, time

        admin = self.make_user("meeting-admin@christuniversity.in", UserProfile.Role.SUPER_ADMIN)
        meeting = Meeting.objects.create(
            title="Cancelled Selection Meeting",
            meeting_date=date(2026, 3, 1),
            start_time=time(10, 0),
            end_time=time(11, 0),
            venue=self.venue.name,
            agenda="Team selection",
            scheduled_by=admin,
            status=Meeting.Status.CANCELLED,
        )
        meeting.teams.add(self.assigned_team)
        self.client.force_login(admin)

        response = self.client.post(reverse("meetings"), {"action": "archive", "meeting_id": meeting.pk})

        self.assertRedirects(response, reverse("meetings"))
        self.assertFalse(Meeting.objects.filter(pk=meeting.pk).exists())
        deleted = Meeting.all_objects.get(pk=meeting.pk)
        self.assertTrue(deleted.is_deleted)
        self.assertEqual(deleted.deleted_by, admin)
        response = self.client.get(reverse("deleted_meetings"))
        self.assertContains(response, meeting.title)
        self.assertContains(response, "Deleted Meeting Details")

    def test_profile_photo_upload_is_saved_and_rendered(self):
        image_bytes = (
            b"GIF87a\x01\x00\x01\x00\x80\x01\x00\x00\x00\x00ccc,\x00\x00\x00\x00"
            b"\x01\x00\x01\x00\x00\x02\x02D\x01\x00;"
        )
        photo = SimpleUploadedFile("profile.gif", image_bytes, content_type="image/gif")
        self.client.force_login(self.assigned_student)

        with tempfile.TemporaryDirectory() as media_root, self.settings(MEDIA_ROOT=media_root):
            response = self.client.post(reverse("my_profile"), {
                "first_name": "Assigned",
                "last_name": "Player",
                "gender": "Male",
                "profile_image": photo,
            })
            self.assertRedirects(response, reverse("my_profile"))
            self.assigned_student.profile.refresh_from_db()
            self.assertTrue(self.assigned_student.profile.profile_image.name.startswith("profile_images/"))
            self.assertTrue(os.path.exists(self.assigned_student.profile.profile_image.path))

            response = self.client.get(reverse("my_profile"))
            self.assertContains(response, f'/media/{self.assigned_student.profile.profile_image.name}')

    def test_requested_management_tables_default_to_one_hundred_rows(self):
        admin = self.make_user("table-admin@christuniversity.in", UserProfile.Role.SUPER_ADMIN)
        self.client.force_login(admin)
        for url_name in ("sports", "teams", "members", "trainers"):
            response = self.client.get(reverse(url_name))
            self.assertContains(response, 'data-default-page-size="100"')

    def test_student_bulk_upload_supports_gender_sport_and_team_type(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Student Name", "Student Email", "Mobile Number", "Reg No", "Department", "Class", "Gender", "Sport", "Team Type", "Team", "Status"])
        sheet.append(["Bulk Player", "bulk.player@christuniversity.in", "9876543210", "B001", "MSDS", "I MSc", "Male", "Cricket", "University", "University Team", "Active"])
        uploaded = BytesIO()
        workbook.save(uploaded)
        uploaded.seek(0)

        created, errors = import_students_from_file(uploaded, self.trainer)

        self.assertEqual(created, 1)
        self.assertEqual(errors, [])
        student = User.objects.get(email="bulk.player@christuniversity.in")
        self.assertEqual(student.profile.gender, "Male")
        self.assertTrue(Membership.objects.filter(user=student, team=self.assigned_team, is_active=True).exists())

    def test_management_bulk_templates_have_dynamic_dropdowns_and_requested_columns(self):
        school = self.assigned_student.profile.school
        if school is None:
            from .models import School
            school = School.objects.create(name="School of Sciences")

        team_book = _management_bulk_workbook("teams")
        team_sheet = team_book["Sports Teams"]
        self.assertEqual([team_sheet.cell(1, column).value for column in range(1, 9)], [
            "Sport", "Team Name", "Team Type", "Gender", "Captain Email", "Vice Captain Email", "Trainer", "Status",
        ])
        self.assertEqual(len(team_sheet.data_validations.dataValidation), 5)
        self.assertEqual(team_book["_Lookups"].sheet_state, "hidden")

        student_book = _management_bulk_workbook("students")
        student_sheet = student_book["Students"]
        self.assertEqual([student_sheet.cell(1, column).value for column in range(1, 12)], [
            "Student Name", "Student Email", "Mobile Number", "Reg No", "School", "Department",
            "Sport", "Team Type", "Gender", "Team", "Status",
        ])
        student_formulas = [validation.formula1 for validation in student_sheet.data_validations.dataValidation]
        self.assertIn("=MgmtSchoolsList", student_formulas)
        self.assertIn("=MgmtGendersList", student_formulas)
        self.assertIn("=INDIRECT($L2)", student_formulas)
        self.assertIn("=INDIRECT($M2)", student_formulas)
        self.assertIn('$G2&"|"&$H2&"|"&$I2', student_sheet["M2"].value)

        trainer_book = _management_bulk_workbook("trainers")
        trainer_sheet = trainer_book["Trainers"]
        self.assertEqual([trainer_sheet.cell(1, column).value for column in range(1, 10)], [
            "Trainer Name", "Trainer Email", "Mobile Number", "Reg No", "Sport", "Team Type", "Gender", "Team", "Status",
        ])
        self.assertIn('$E2&"|"&$F2&"|"&$G2', trainer_sheet["K2"].value)
        self.assertIn("=MgmtStatusesList", [validation.formula1 for validation in trainer_sheet.data_validations.dataValidation])

    def test_all_applicable_bulk_templates_have_active_inactive_status_dropdown(self):
        for kind in ("sports", "venues", "schools"):
            workbook = _simple_status_bulk_workbook(kind)
            sheet = workbook["Sample"]
            self.assertIn("Status", [cell.value for cell in sheet[1]])
            self.assertEqual(
                [validation.formula1 for validation in sheet.data_validations.dataValidation],
                ["=SimpleStatusesList"],
            )
            lookup_values = [workbook["_Lookups"].cell(row, 1).value for row in (2, 3)]
            self.assertEqual(lookup_values, ["Active", "Inactive"])

        for kind in ("teams", "students", "trainers"):
            workbook = _management_bulk_workbook(kind)
            sheet = workbook[{"teams": "Sports Teams", "students": "Students", "trainers": "Trainers"}[kind]]
            self.assertIn("Status", [cell.value for cell in sheet[1]])
            self.assertIn("=MgmtStatusesList", [validation.formula1 for validation in sheet.data_validations.dataValidation])

    def test_bulk_imports_reject_nonstandard_status_values(self):
        cases = (
            (import_sports_from_file, ["Sport Name", "Description", "Status"], ["Status Sport", "", "Deactivated"]),
            (import_venues_from_file, ["Venue Name", "Location", "Status"], ["Status Venue", "", "Yes"]),
            (import_schools_from_file, ["School Name", "Description", "Status"], ["Status School", "", "Disabled"]),
        )
        for importer, headers, values in cases:
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(headers)
            sheet.append(values)
            uploaded = BytesIO()
            workbook.save(uploaded)
            uploaded.seek(0)
            created, errors = importer(uploaded)
            self.assertEqual(created, 0)
            self.assertEqual(errors, ["Row 2: Status must be Active or Inactive."])

    def test_inactive_status_is_saved_for_sport_and_trainer(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Sport Name", "Description", "Status"])
        sheet.append(["Inactive Sport", "", "Inactive"])
        uploaded = BytesIO()
        workbook.save(uploaded)
        uploaded.seek(0)
        created, errors = import_sports_from_file(uploaded)
        self.assertEqual((created, errors), (1, []))
        self.assertFalse(Sport.objects.get(name="Inactive Sport").is_active)

        trainer_book = Workbook()
        trainer_sheet = trainer_book.active
        trainer_sheet.append(["Trainer Name", "Trainer Email", "Mobile Number", "Reg No", "Status"])
        trainer_sheet.append(["Inactive Coach", "inactive.coach@christuniversity.in", "", "IC01", "Inactive"])
        trainer_upload = BytesIO()
        trainer_book.save(trainer_upload)
        trainer_upload.seek(0)
        created, errors = import_trainers_from_file(trainer_upload)
        self.assertEqual((created, errors), (1, []))
        trainer = User.objects.get(email="inactive.coach@christuniversity.in")
        self.assertFalse(trainer.is_active)
        self.assertEqual(trainer.profile.account_status, UserProfile.AccountStatus.DEACTIVATED)

    def test_team_bulk_upload_validates_dropdown_values_and_assigns_trainer(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Sport", "Team Name", "Team Type", "Gender", "Captain Email", "Vice Captain Email", "Trainer", "Status"])
        sheet.append(["Cricket", "New Team", "Reserve", "Female", "", "", f"{self.trainer.email}", "Active"])
        sheet.append(["Unknown Sport", "Bad Team", "University", "Male", "", "", "", "Active"])
        sheet.append(["Cricket", "Bad Type", "Invalid", "Male", "", "", "", "Active"])
        uploaded = BytesIO()
        workbook.save(uploaded)
        uploaded.seek(0)

        created, errors = import_teams_from_file(uploaded)

        self.assertEqual(created, 1)
        self.assertEqual(len(errors), 2)
        self.assertEqual(Team.objects.get(name="New Team").coordinator, self.trainer)
        self.assertIn("Unknown Sport", errors[0])
        self.assertIn("Team Type", errors[1])

    def test_student_team_dropdown_combination_is_enforced(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Student Name", "Student Email", "Mobile Number", "Reg No", "School", "Department", "Sport", "Team Type", "Team", "Gender", "Status"])
        sheet.append(["Valid Student", "valid.bulk@christuniversity.in", "", "V001", "", "MSDS", "Cricket", "University", "University Team", "Male", "Active"])
        sheet.append(["Invalid Student", "invalid.bulk@christuniversity.in", "", "V002", "", "MSDS", "Cricket", "University", "University Team", "Female", "Active"])
        uploaded = BytesIO()
        workbook.save(uploaded)
        uploaded.seek(0)

        created, errors = import_students_from_file(uploaded, self.trainer)

        self.assertEqual(created, 1)
        self.assertEqual(len(errors), 1)
        self.assertTrue(Membership.objects.filter(user__email="valid.bulk@christuniversity.in", team=self.assigned_team).exists())
        self.assertFalse(User.objects.filter(email="invalid.bulk@christuniversity.in").exists())
        self.assertIn("No active Team matched", errors[0])

    def test_trainer_bulk_upload_supports_exact_team_hierarchy_and_multiple_rows(self):
        second_team = Team.objects.create(
            sport=self.other_sport,
            name="Second Team",
            team_type=Team.TeamType.RECREATIONAL,
            gender=Team.TeamGender.FEMALE,
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Trainer Name", "Trainer Email", "Mobile Number", "Reg No", "Sport", "Team Type", "Team", "Gender"])
        sheet.append(["Bulk Coach", "bulk.coach@christuniversity.in", "9876543212", "TC01", "Cricket", "University", "University Team", "Male"])
        sheet.append(["Bulk Coach", "bulk.coach@christuniversity.in", "9876543212", "TC01", "Football", "Recreational", "Second Team", "Female"])
        uploaded = BytesIO()
        workbook.save(uploaded)
        uploaded.seek(0)

        created, errors = import_trainers_from_file(uploaded)

        self.assertEqual(created, 2)
        self.assertEqual(errors, [])
        trainer = User.objects.get(email="bulk.coach@christuniversity.in")
        self.assigned_team.refresh_from_db()
        second_team.refresh_from_db()
        self.assertEqual(self.assigned_team.coordinator, trainer)
        self.assertEqual(second_team.coordinator, trainer)

    @patch("core.views.create_notifications")
    def test_session_bulk_upload_matches_all_four_team_fields(self, mocked_notifications):
        target_team = Team.objects.create(
            sport=self.sport,
            name="Shared Team",
            team_type=Team.TeamType.RESERVE,
            gender=Team.TeamGender.FEMALE,
            coordinator=self.trainer,
        )
        Team.objects.create(
            sport=self.sport,
            name="Shared Team",
            team_type=Team.TeamType.UNIVERSITY,
            gender=Team.TeamGender.FEMALE,
            coordinator=self.trainer,
        )
        Team.objects.create(
            sport=self.sport,
            name="Shared Team",
            team_type=Team.TeamType.RESERVE,
            gender=Team.TeamGender.MALE,
            coordinator=self.trainer,
        )
        Team.objects.create(
            sport=self.other_sport,
            name="Shared Team",
            team_type=Team.TeamType.RESERVE,
            gender=Team.TeamGender.FEMALE,
            coordinator=self.trainer,
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Sport", "Team Name", "Team Type", "Gender", "Title", "Start Date", "End Date", "From Time", "To Time", "Schedule", "Full Day", "Venue", "Other Venue", "Notes"])
        sheet.append(["Cricket", "Shared Team", "Reserve", "Female", "Exact Team Practice", "2026-11-01", "2026-11-01", "06:30", "08:30", "Morning", "No", "Ground A", "", "Drills"])
        uploaded = BytesIO()
        workbook.save(uploaded)
        uploaded.seek(0)

        created, errors = import_sessions_from_file(uploaded, self.trainer)

        self.assertEqual(created, 1)
        self.assertEqual(errors, [])
        self.assertEqual(Session.objects.get(title="Exact Team Practice").team, target_team)
        mocked_notifications.assert_called_once()

    def test_session_bulk_upload_reports_missing_and_invalid_team_details(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Sport", "Team Name", "Team Type", "Gender", "Title", "Start Date", "End Date", "From Time", "To Time", "Schedule", "Full Day", "Venue", "Other Venue", "Notes"])
        sheet.append(["Cricket", "", "University", "Male", "Missing Team", "2026-11-01", "2026-11-01", "06:30", "08:30", "Morning", "No", "Ground A", "", ""])
        sheet.append(["Cricket", "University Team", "Invalid Type", "Male", "Invalid Type", "2026-11-01", "2026-11-01", "06:30", "08:30", "Morning", "No", "Ground A", "", ""])
        sheet.append(["Cricket", "University Team", "University", "Unknown", "Invalid Gender", "2026-11-01", "2026-11-01", "06:30", "08:30", "Morning", "No", "Ground A", "", ""])
        uploaded = BytesIO()
        workbook.save(uploaded)
        uploaded.seek(0)

        created, errors = import_sessions_from_file(uploaded, self.trainer)

        self.assertEqual(created, 0)
        self.assertEqual(len(errors), 3)
        self.assertIn("Team Name required", errors[0])
        self.assertIn("Team Type 'Invalid Type' is invalid", errors[1])
        self.assertIn("Gender 'Unknown' is invalid", errors[2])

    def test_session_template_has_dynamic_dependent_dropdowns_and_date_format(self):
        Team.objects.create(
            sport=self.sport,
            name="Shared Team",
            team_type=Team.TeamType.RESERVE,
            gender=Team.TeamGender.FEMALE,
            coordinator=self.trainer,
        )

        workbook = build_sessions_bulk_upload_workbook()
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        reloaded = load_workbook(output, data_only=False)
        sheet = reloaded["Practice Sessions"]

        self.assertEqual(reloaded["_Lookups"].sheet_state, "hidden")
        self.assertEqual([sheet.cell(1, index).value for index in range(1, 5)], ["Sport", "Team Type", "Team Name", "Gender"])
        self.assertEqual(sheet["F2"].number_format, "DD/MM/YYYY")
        self.assertEqual(sheet["G2"].number_format, "DD/MM/YYYY")
        self.assertIn("VLOOKUP($A2", sheet["O2"].value)
        formulas = [validation.formula1 for validation in sheet.data_validations.dataValidation]
        self.assertIn("=SportsList", formulas)
        self.assertIn("=SchedulesList", formulas)
        self.assertIn("=VenuesList", formulas)
        self.assertIn("=INDIRECT($O2)", formulas)
        self.assertIn("=INDIRECT($P2)", formulas)
        self.assertIn("=INDIRECT($Q2)", formulas)

        output.seek(0)
        parsed_rows = read_bulk_upload_rows(output)
        self.assertEqual(len(parsed_rows), 1)
        self.assertNotIn("teamtyperange", parsed_rows[0])

    @patch("core.views.create_notifications")
    def test_session_bulk_upload_ignores_blank_formula_rows_and_keeps_excel_row_numbers(self, mocked_notifications):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Sport", "Team Type", "Team Name", "Gender", "Title", "Start Date", "End Date",
            "From Time", "To Time", "Schedule", "Full Day", "Venue", "Other Venue", "Notes",
            "_TeamTypeRange", "_TeamRange", "_GenderRange",
        ])
        valid = ["Cricket", "University", "University Team", "Male", "Practice", "01/11/2026", "01/11/2026", "06:30", "08:30", "Morning", "No", "Ground A", "", ""]
        sheet.append([*valid, "DependentList1", "DependentList2", "DependentList3"])
        sheet.append([None] * 14 + ["EmptyList", "EmptyList", "EmptyList"])
        sheet.append([*valid, "DependentList1", "DependentList2", "DependentList3"])
        sheet.append(["Cricket"] + [None] * 13 + ["DependentList1", "EmptyList", "EmptyList"])
        for _ in range(496):
            sheet.append([None] * 14 + ["EmptyList", "EmptyList", "EmptyList"])
        uploaded = BytesIO()
        workbook.save(uploaded)
        uploaded.seek(0)

        created, errors = import_sessions_from_file(uploaded, self.trainer)

        self.assertEqual(created, 2)
        self.assertEqual(len(errors), 1)
        self.assertTrue(errors[0].startswith("Row 5:"))
        self.assertIn("Team Name", errors[0])
        self.assertEqual(mocked_notifications.call_count, 2)

    @patch("core.views.create_notifications")
    def test_session_bulk_upload_accepts_excel_dates_and_ddmmyyyy_rows(self, mocked_notifications):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Sport", "Team Type", "Team Name", "Gender", "Title", "Start Date", "End Date", "From Time", "To Time", "Schedule", "Full Day", "Venue", "Other Venue", "Notes"])
        sheet.append(["Cricket", "University", "University Team", "Male", "Excel Date", timezone.localdate(), timezone.localdate(), "06:30", "08:30", "Morning", "No", "Ground A", "", ""])
        sheet.append(["Cricket", "University", "University Team", "Male", "Text Date", "02/11/2026", "03/11/2026", "06:30", "08:30", "Evening", "No", "Ground A", "", ""])
        uploaded = BytesIO()
        workbook.save(uploaded)
        uploaded.seek(0)

        created, errors = import_sessions_from_file(uploaded, self.trainer)

        self.assertEqual(created, 2)
        self.assertEqual(errors, [])
        self.assertEqual(Session.objects.get(title="Text Date").schedule_slot, Session.ScheduleSlot.EVENING)
        self.assertEqual(mocked_notifications.call_count, 2)

    def test_session_bulk_upload_rejects_stale_schedule_venue_and_team(self):
        stale_team = Team.objects.create(
            sport=self.sport,
            name="Old Team",
            team_type=Team.TeamType.RESERVE,
            gender=Team.TeamGender.FEMALE,
            coordinator=self.trainer,
            is_active=False,
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Sport", "Team Type", "Team Name", "Gender", "Title", "Start Date", "End Date", "From Time", "To Time", "Schedule", "Full Day", "Venue", "Other Venue", "Notes"])
        sheet.append(["Cricket", "Reserve", stale_team.name, "Female", "Stale Team", "01/11/2026", "01/11/2026", "06:30", "08:30", "Morning", "No", "Ground A", "", ""])
        sheet.append(["Cricket", "University", "University Team", "Male", "Bad Schedule", "01/11/2026", "01/11/2026", "06:30", "08:30", "Night", "No", "Ground A", "", ""])
        sheet.append(["Cricket", "University", "University Team", "Male", "Stale Venue", "01/11/2026", "01/11/2026", "06:30", "08:30", "Morning", "No", "Deleted Ground", "", ""])
        uploaded = BytesIO()
        workbook.save(uploaded)
        uploaded.seek(0)

        created, errors = import_sessions_from_file(uploaded, self.trainer)

        self.assertEqual(created, 0)
        self.assertEqual(len(errors), 3)
        self.assertIn("No team matched", errors[0])
        self.assertIn("Schedule 'Night' is invalid", errors[1])
        self.assertIn("Venue 'Deleted Ground' was not found or is inactive", errors[2])
